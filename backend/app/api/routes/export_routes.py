"""Export endpoints: Excel and PDF reports - FDWH version."""

import io
from datetime import datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.audit import audit
from app.core.constants import EXCEL_MAX_DATA_ROWS, EXPORT_DEFAULT_DAYS
from app.db.database import get_db
from app.db import queries
from app.auth.auth import get_current_user
from app.models.models import User, parse_component_id, fdwh_unit, FDWH_MEASUREMENT_TYPES

router = APIRouter(prefix="/api/export", tags=["Export"])

MTYPE_LABELS = {
    "P": "Wirkleistung", "Q": "Blindleistung", "S": "Scheinleistung",
    "U": "Spannung", "I": "Strom",
}


def _default_time_range():
    now = datetime.now()
    return now - timedelta(days=EXPORT_DEFAULT_DAYS), now


def _comp_name(db: Session, anr: str, fnr: str) -> str:
    row = db.execute(
        text(queries.component_name()), {"anr": anr, "fnr": fnr}
    ).fetchone()
    return row[0] if row else ""


@router.get("/{component_id}/excel")
def export_excel(
    component_id: str,
    start: Optional[datetime] = Query(None),
    end: Optional[datetime] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    anr, fnr = parse_component_id(component_id)
    comp_name = _comp_name(db, anr, fnr)
    if not comp_name:
        raise HTTPException(status_code=404, detail="Komponente nicht gefunden")
    audit("export.excel", username=current_user.username, component_id=component_id)

    if not start or not end:
        start, end = _default_time_range()

    wb = Workbook()

    # --- Sheet 1: Overview ---
    ws = wb.active
    ws.title = "Übersicht"

    header_font = Font(bold=True, size=14)
    sub_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_text = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    ws.merge_cells('A1:D1')
    ws['A1'] = f"Betriebsmittel-Bericht: {comp_name}"
    ws['A1'].font = header_font

    ws['A3'] = "Anlagennummer:"
    ws['B3'] = anr
    ws['A4'] = "Feldnummer:"
    ws['B4'] = fnr
    ws['A5'] = "Zeitraum:"
    ws['B5'] = f"{start.strftime('%d.%m.%Y %H:%M')} - {end.strftime('%d.%m.%Y %H:%M')}"

    for cell in ['A3', 'A4', 'A5']:
        ws[cell].font = Font(bold=True)

    # Statistics table
    ws['A7'] = "Statistik"
    ws['A7'].font = sub_font

    stats_header = ['Messgröße', 'Einheit', 'Anzahl', 'Mittelwert', 'Min', 'Max', 'Std.Abw.']
    for col, h in enumerate(stats_header, 1):
        cell = ws.cell(row=8, column=col, value=h)
        cell.font = header_text
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    s_str = start.strftime("%Y-%m-%d %H:%M:%S")
    e_str = end.strftime("%Y-%m-%d %H:%M:%S")

    row_num = 9
    for mtype in FDWH_MEASUREMENT_TYPES:
        unit = fdwh_unit(mtype)
        stats_row = db.execute(
            text(queries.statistics(mtype)),
            {"anr": anr, "fnr": fnr, "start": s_str, "end": e_str},
        ).fetchone()

        cnt = stats_row[0] if stats_row and stats_row[0] else 0
        mean_val = stats_row[1] if stats_row and stats_row[1] else 0
        min_val = stats_row[2] if stats_row and stats_row[2] is not None else 0
        max_val = stats_row[3] if stats_row and stats_row[3] is not None else 0
        std_val = stats_row[4] if stats_row and stats_row[4] else ""

        values = [
            MTYPE_LABELS.get(mtype, mtype), unit, cnt,
            round(mean_val, 3), round(min_val, 3), round(max_val, 3),
            round(std_val, 3) if std_val != "" else "",
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col, value=v)
            cell.border = thin_border
            if col >= 3:
                cell.alignment = Alignment(horizontal='right')
        row_num += 1

    for col in range(1, 8):
        ws.column_dimensions[ws.cell(row=8, column=col).column_letter].width = 16

    # --- Sheet 2: Timeseries Data ---
    ws2 = wb.create_sheet("Zeitreihen")

    header = ['Zeitstempel'] + [f"{MTYPE_LABELS.get(mt, mt)} ({fdwh_unit(mt)})" for mt in FDWH_MEASUREMENT_TYPES]
    for col, h in enumerate(header, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_text
        cell.fill = header_fill
        cell.border = thin_border

    rows_data = db.execute(
        text(queries.export_timeseries()),
        {"anr": anr, "fnr": fnr, "s": s_str, "e": e_str},
    ).fetchall()

    # Each row: (LOKALZEIT, MW, BMW, S, UUW, strom)
    row_num = 2
    for r in rows_data:
        ws2.cell(row=row_num, column=1, value=str(r[0])).border = thin_border
        for col_idx in range(1, 6):
            val = r[col_idx]
            cell = ws2.cell(row=row_num, column=col_idx + 1,
                            value=round(val, 3) if val is not None else "")
            cell.border = thin_border
        row_num += 1
        if row_num > EXCEL_MAX_DATA_ROWS + 1:
            break

    ws2.column_dimensions['A'].width = 22
    for col in range(2, len(FDWH_MEASUREMENT_TYPES) + 2):
        ws2.column_dimensions[ws2.cell(row=1, column=col).column_letter].width = 16

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"bericht_{comp_name.replace(' ', '_')}_{start.strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{component_id}/pdf")
def export_pdf(
    component_id: str,
    start: Optional[datetime] = Query(None),
    end: Optional[datetime] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    anr, fnr = parse_component_id(component_id)
    comp_name = _comp_name(db, anr, fnr)
    if not comp_name:
        raise HTTPException(status_code=404, detail="Komponente nicht gefunden")
    audit("export.pdf", username=current_user.username, component_id=component_id)

    if not start or not end:
        start, end = _default_time_range()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f"Betriebsmittel-Bericht: {comp_name}", styles['Title']))
    elements.append(Spacer(1, 12))

    info_data = [
        ["Anlagennummer:", anr],
        ["Feldnummer:", fnr],
        ["Zeitraum:", f"{start.strftime('%d.%m.%Y %H:%M')} - {end.strftime('%d.%m.%Y %H:%M')}"],
    ]

    info_table = Table(info_data, colWidths=[5 * cm, 10 * cm])
    info_table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 20))

    elements.append(Paragraph("Statistik", styles['Heading2']))
    elements.append(Spacer(1, 8))

    s_str = start.strftime("%Y-%m-%d %H:%M:%S")
    e_str = end.strftime("%Y-%m-%d %H:%M:%S")

    stats_data = [["Messgröße", "Einheit", "Anzahl", "Mittelwert", "Min", "Max"]]
    for mtype in FDWH_MEASUREMENT_TYPES:
        unit = fdwh_unit(mtype)
        stats_row = db.execute(
            text(queries.statistics(mtype)),
            {"anr": anr, "fnr": fnr, "start": s_str, "end": e_str},
        ).fetchone()

        cnt = stats_row[0] if stats_row and stats_row[0] else 0
        mean_val = stats_row[1] if stats_row and stats_row[1] else 0
        min_val = stats_row[2] if stats_row and stats_row[2] is not None else 0
        max_val = stats_row[3] if stats_row and stats_row[3] is not None else 0

        stats_data.append([
            MTYPE_LABELS.get(mtype, mtype), unit, str(cnt),
            f"{mean_val:.2f}", f"{min_val:.2f}", f"{max_val:.2f}",
        ])

    stats_table = Table(stats_data, colWidths=[3.5 * cm, 2 * cm, 2 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm])
    stats_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E79')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (2, 0), (-1, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F4F8')]),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(stats_table)

    doc.build(elements)
    buffer.seek(0)

    filename = f"bericht_{comp_name.replace(' ', '_')}_{start.strftime('%Y%m%d')}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
